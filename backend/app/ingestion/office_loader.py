"""Microsoft Office ingestion: Word, Excel and PowerPoint."""

from __future__ import annotations

from pathlib import Path

from ..models.schemas import ExtractedTable


def load_docx(path: Path) -> tuple[str, list[ExtractedTable], dict]:
    """Extract paragraphs and tables from a Word .docx file."""
    from docx import Document

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    tables: list[ExtractedTable] = []
    for tbl in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]
        rows = [r for r in rows if any(c for c in r)]
        if rows:
            tables.append(ExtractedTable(rows=rows))

    # Fold table content into the text stream too so RAG/LLM see it.
    table_text = "\n".join(
        "\t".join(r) for t in tables for r in t.rows
    )
    text = "\n".join(paragraphs)
    if table_text:
        text = f"{text}\n\n{table_text}"
    return text, tables, {"paragraphs": len(paragraphs), "tables": len(tables)}


def load_xlsx(path: Path) -> tuple[str, list[ExtractedTable], dict]:
    """Extract every worksheet of an Excel workbook as a table + text."""
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    tables: list[ExtractedTable] = []
    text_parts: list[str] = []

    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(cells)
        if rows:
            tables.append(ExtractedTable(caption=ws.title, rows=rows))
            text_parts.append(f"# Sheet: {ws.title}")
            text_parts.extend("\t".join(r) for r in rows)

    wb.close()
    return "\n".join(text_parts), tables, {"sheets": len(tables)}


def load_pptx(path: Path) -> tuple[str, dict]:
    """Extract text from every slide of a PowerPoint deck."""
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    slide_count = 0
    for idx, slide in enumerate(prs.slides, start=1):
        slide_count = idx
        slide_lines = [f"# Slide {idx}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        slide_lines.append(line)
        if len(slide_lines) > 1:
            parts.append("\n".join(slide_lines))
    return "\n\n".join(parts), {"slides": slide_count}
