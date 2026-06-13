"""Tests for document ingestion."""

from __future__ import annotations

from docx import Document
from openpyxl import Workbook

from app.ingestion import load_document


def test_load_text(tmp_path):
    f = tmp_path / "note.txt"
    f.write_text("Start pump P-101 and verify discharge pressure.", encoding="utf-8")
    doc = load_document(f)
    assert "P-101" in doc.text
    assert doc.filename == "note.txt"


def test_load_docx(tmp_path):
    f = tmp_path / "proc.docx"
    d = Document()
    d.add_paragraph("Commissioning procedure for pump P-101.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Tag"
    table.rows[0].cells[1].text = "20-P-101A"
    d.save(str(f))

    doc = load_document(f)
    assert "Commissioning procedure" in doc.text
    assert doc.tables and "20-P-101A" in doc.text


def test_load_xlsx(tmp_path):
    f = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Tags"
    ws.append(["Tag", "Service"])
    ws.append(["20-P-101A", "Feed Pump"])
    wb.save(str(f))

    doc = load_document(f)
    assert "20-P-101A" in doc.text
    assert doc.tables and doc.tables[0].caption == "Tags"
